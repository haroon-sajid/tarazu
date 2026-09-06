"""What client data must never be able to do to the firm that audits it.

Every value in a case is whatever the client's bookkeeper typed. A party name
is also a cell in the auditor's Excel report, a filename is also a storage
path and a download header, and both come from outside. These tests pin the
three places that matter:

- **Formula injection.** openpyxl stores a string beginning with `=` as a
  formula unless told otherwise. A ledger party `=HYPERLINK(...)` must reach
  the report and the analytics workbook as text.
- **Path traversal.** `../../etc/passwd.csv` is a ledger to the reader and an
  attack to the store. Only the basename survives, and the local store refuses
  any path that resolves outside its root.
- **Header safety.** A filename with quotes or a newline must not be able to
  break the `Content-Disposition` header the download route writes.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.api.files import MAX_FILENAME_LENGTH, safe_filename, suffix_of
from app.core.sqlite_store import LocalDocumentStore, SqliteCaseRepository
from app.modules.analytics import service as analytics
from app.modules.reports import service as reports
from app.shared.schemas import ReviewDecision
from tests.conftest import DEMO_ORG_ID, DEMO_USER

GENERATED_AT = datetime(2026, 9, 6, 10, 0, tzinfo=timezone.utc)

#: What a hostile bookkeeper can put in a cell.
FORMULAS = (
    '=HYPERLINK("http://evil.example/x","Open me")',
    "=1+1",
    "=cmd|' /C calc'!A0",
    "+SUM(A1:A9)",
    "-2+3",
    "@SUM(1,2)",
)


def _text_cells(workbook_bytes: bytes) -> dict[str, list[tuple[str, str]]]:
    """Every string cell in the workbook, as (value, data_type), by sheet."""
    workbook = load_workbook(io.BytesIO(workbook_bytes))
    found: dict[str, list[tuple[str, str]]] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    found.setdefault(sheet.title, []).append((cell.value, cell.data_type))
    return found


# --------------------------------------------------------------------------- #
# 1. Formula injection
# --------------------------------------------------------------------------- #


def test_a_party_name_that_looks_like_a_formula_stays_text_in_the_report(
    repository: SqliteCaseRepository, seeded_case: str
) -> None:
    items = repository.list_review_items(DEMO_ORG_ID, seeded_case)
    # Decided items are what the report lists, so the hostile name goes on one.
    first = items[0].model_copy(
        update={
            "decision": ReviewDecision.APPROVED,
            "decided_by": DEMO_USER.user_id,
            "decided_at": GENERATED_AT,
            "ledger_entry": items[0].ledger_entry.model_copy(
                update={"party_name": FORMULAS[0], "description": FORMULAS[2]}
            ),
        }
    )
    files = reports.generate_report(
        repository.get_case(DEMO_ORG_ID, seeded_case),
        [first, *items[1:]],
        repository.list_audit(DEMO_ORG_ID, seeded_case),
        repository.get_benford(DEMO_ORG_ID, seeded_case),
        report_id="RPT-hardening",
        generated_by=DEMO_USER.user_id,
        generated_at=GENERATED_AT,
    )

    cells = [cell for sheet in _text_cells(files.excel).values() for cell in sheet]
    hostile = [cell for cell in cells if cell[0].startswith(("=", "+", "@"))]
    assert hostile, "the hostile name should appear in the workbook at all"
    assert all(data_type == "s" for _value, data_type in hostile), hostile
    assert all(data_type != "f" for _value, data_type in cells)


def _csv(header: list[str], rows: list[list[object]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def test_a_customer_or_product_that_looks_like_a_formula_stays_text_in_the_analytics_export() -> None:
    rows = _csv(
        ["Date", "Customer", "Product", "Region", "Amount"],
        [[f"0{day}/06/2026", FORMULAS[0], FORMULAS[2], FORMULAS[3], 1000 * day] for day in range(1, 8)],
    )
    records, report = analytics.read_sales_export("SLS-H", "sales.csv", rows)
    result = analytics.analyze_sales(records, [report])
    workbook = analytics.export_workbook(result)

    cells = [cell for sheet in _text_cells(workbook).values() for cell in sheet]
    assert any(value.startswith("=") for value, _ in cells)
    assert all(data_type != "f" for _value, data_type in cells), [
        cell for cell in cells if cell[1] == "f"
    ]


# --------------------------------------------------------------------------- #
# 2. Filenames
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("ledger.xlsx", "ledger.xlsx"),
        ("../../etc/passwd.csv", "passwd.csv"),
        ("..\\..\\windows\\evil.csv", "evil.csv"),
        ("C:\\Users\\x\\Desktop\\June ledger.xlsx", "June ledger.xlsx"),
        ("/tmp/statement.pdf", "statement.pdf"),
        ("  spaced.csv  ", "spaced.csv"),
        ("..", "unnamed"),
        ("", "unnamed"),
        (None, "unnamed"),
        ('quote"d\r\nname.csv', "quote_d_name.csv"),
        ("tab\tin\x00name.pdf", "tab_in_name.pdf"),
        ("what?.csv", "what_.csv"),
        ("....hidden.csv", "hidden.csv"),
    ],
)
def test_safe_filename_keeps_the_basename_and_nothing_dangerous(raw, expected) -> None:
    assert safe_filename(raw) == expected


def test_an_overlong_filename_is_cut_but_keeps_its_extension() -> None:
    name = safe_filename("x" * 400 + ".xlsx")
    assert len(name) <= MAX_FILENAME_LENGTH
    assert name.endswith(".xlsx")
    assert suffix_of(name) == ".xlsx"

    no_extension = safe_filename("y" * 400)
    assert len(no_extension) == MAX_FILENAME_LENGTH


def test_suffix_of_is_case_insensitive_and_tolerates_odd_names() -> None:
    assert suffix_of("Ledger.CSV") == ".csv"
    assert suffix_of("archive.tar.gz") == ".gz"
    assert suffix_of("noext") == ""
    assert suffix_of(None) == ""
    assert suffix_of(".csv") == ".csv"


# --------------------------------------------------------------------------- #
# 3. The local document store
# --------------------------------------------------------------------------- #


def test_the_local_store_refuses_a_path_outside_its_root(tmp_path) -> None:
    store = LocalDocumentStore(tmp_path / "store")
    (tmp_path / "store2").mkdir()

    for escape in ("../escape.txt", "CASE/../../escape.txt", "../store2/x.txt", "."):
        with pytest.raises(ValueError, match="escapes"):
            store.put(escape, b"x", "text/plain")
        with pytest.raises(ValueError, match="escapes"):
            store.get(escape)

    # Nothing landed outside.
    assert not (tmp_path / "escape.txt").exists()
    assert not (tmp_path / "store2" / "x.txt").exists()


def test_the_local_store_round_trips_a_nested_path(tmp_path) -> None:
    store = LocalDocumentStore(tmp_path / "store")
    assert store.put("CASE-1/DOC-1/ledger.xlsx", b"bytes", "x") == "CASE-1/DOC-1/ledger.xlsx"
    assert store.get("CASE-1/DOC-1/ledger.xlsx") == b"bytes"


# --------------------------------------------------------------------------- #
# 4. Through the routes
# --------------------------------------------------------------------------- #


def _upload(client, ledger_name: str):
    from tests.test_pipeline import a_ledger, a_pdf

    return client.post(
        "/v1/upload",
        files=[
            ("bank_statement", ("statement.pdf", io.BytesIO(a_pdf()))),
            ("ledger", (ledger_name, io.BytesIO(a_ledger()))),
            ("invoices", ("invoice.pdf", io.BytesIO(a_pdf("INVOICE")))),
        ],
    )


def test_a_traversal_filename_is_stored_under_its_basename(client, storage, demo_mode) -> None:
    response = _upload(client, "..\\..\\..\\evil-ledger.xlsx")
    assert response.status_code == 201, response.text
    body = response.json()
    ledger = next(doc for doc in body["documents"] if doc["document_type"] == "ledger")
    assert ledger["filename"] == "evil-ledger.xlsx"
    assert ".." not in ledger["storage_path"]
    # And the bytes are where the path says, inside the store.
    assert storage.get(ledger["storage_path"])


def test_a_filename_with_quotes_and_a_newline_still_downloads(client, demo_mode) -> None:
    response = _upload(client, 'led"ger\r\nX-Injected: yes.xlsx')
    assert response.status_code == 201, response.text
    ledger = next(
        doc for doc in response.json()["documents"] if doc["document_type"] == "ledger"
    )
    assert "\n" not in ledger["filename"] and '"' not in ledger["filename"]

    download = client.get(f"/v1/documents/{ledger['document_id']}/file")
    assert download.status_code == 200
    assert "x-injected" not in {name.lower() for name in download.headers}
    assert "\n" not in download.headers.get("content-disposition", "")


def test_a_hostile_party_name_survives_the_whole_pipeline_as_text(
    client, repository: SqliteCaseRepository, demo_mode
) -> None:
    """From the ledger cell to the review queue to the workbook, verbatim and inert.

    A CSV, because a workbook written by a library stores `=...` as a formula
    with no cached value, which reads back as nothing — a fact about the
    fixture, not the product. A bookkeeper's saved workbook carries the value.
    """
    from tests.test_pipeline import a_pdf

    ledger = _csv(
        ["Date", "Party Name", "Amount", "Particulars"],
        [["02/06/2026", FORMULAS[0], 284000, FORMULAS[2]]],
    )
    response = client.post(
        "/v1/upload",
        files=[
            ("bank_statement", ("statement.pdf", io.BytesIO(a_pdf()))),
            ("ledger", ("ledger.csv", io.BytesIO(ledger))),
            ("invoices", ("invoice.pdf", io.BytesIO(a_pdf("INVOICE")))),
        ],
    )
    assert response.status_code == 201, response.text
    case_id = response.json()["case_id"]

    items = client.get(f"/v1/review-items?case_id={case_id}").json()["items"]
    assert items[0]["ledger_entry"]["party_name"] == FORMULAS[0]

    approved = client.post(
        f"/v1/review-items/{items[0]['review_item_id']}/approve", json={"note": "seen"}
    )
    assert approved.status_code == 200, approved.text

    report = client.post("/v1/reports", json={"case_id": case_id})
    assert report.status_code == 201, report.text
    excel = client.get(report.json()["downloads"]["excel"])
    assert excel.status_code == 200
    cells = [cell for sheet in _text_cells(excel.content).values() for cell in sheet]
    assert any(value == FORMULAS[0] for value, _ in cells)
    assert all(data_type != "f" for _value, data_type in cells)
